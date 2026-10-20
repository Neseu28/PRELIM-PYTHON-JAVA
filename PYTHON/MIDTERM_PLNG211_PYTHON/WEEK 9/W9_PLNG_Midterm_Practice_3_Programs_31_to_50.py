# WEEK 9
# PLNG - Midterm Practice #3


def program_31():
    import pymysql.cursors
    connection = pymysql.connect(host="localhost", user="root", password="",
        db="students", charset="utf8mb4",
        cursorclass=pymysql.cursors.DictCursor)
    try:
        with connection.cursor() as cursor:
            sql = "SELECT `id`, `firstname`, `lastname` FROM `users`"
            cursor.execute(sql)
            for row in cursor.fetchall():
                print("firstname : " + str(row["firstname"]))
                print("lastname : " + str(row["lastname"]))
    finally:
        connection.close()


def program_32():
    from tkinter import Tk
    window = Tk()
    window.title("Hello Python")
    window.geometry("300x200+10+20")
    window.mainloop()


def program_33():
    from tkinter import Tk, Frame, Label, Entry
    window = Tk()
    window.title("mrhuseyin.medium.com")
    window.geometry("400x300")
    application = Frame(window)
    application.grid()
    label = Label(application, text="Enter your name")
    label.grid(padx=110, pady=10)
    entry = Entry(application, bd=2)
    entry.grid(padx=110, pady=3)
    window.mainloop()


def program_34():
    from tkinter import Tk, Frame, Listbox
    window = Tk()
    window.title("mrhuseyin.medium.com")
    window.geometry("400x300")
    application = Frame(window)
    application.grid()
    listbox = Listbox(application)
    for index, item in enumerate(["Python", "C#", "JAVA", "JAVASCRIPT"], start=1):
        listbox.insert(index, item)
    listbox.grid(padx=110, pady=10)
    window.mainloop()


def program_35():
    # Same registration-form idea as the PDF, using excel.xlsx in this folder
    # instead of a machine-specific Windows path.
    from tkinter import Tk, Label, Entry, Button, END, messagebox
    from openpyxl import Workbook, load_workbook
    from pathlib import Path

    excel_file = Path("excel.xlsx")
    if excel_file.exists():
        wb = load_workbook(excel_file)
        sheet = wb.active
    else:
        wb = Workbook()
        sheet = wb.active
        sheet.append(["Name", "Course", "Semester", "Form Number",
                      "Contact Number", "Email id", "Address"])
        wb.save(excel_file)

    root = Tk()
    root.configure(background="light green")
    root.title("registration form")
    root.geometry("600x350")
    labels = ["Name", "Course", "Semester", "Form No.",
              "Contact No.", "Email id", "Address"]
    fields = []
    Label(root, text="Form", bg="light green").grid(row=0, column=1)
    for row, text in enumerate(labels, start=1):
        Label(root, text=text, bg="light green").grid(
            row=row, column=0, padx=10, pady=3, sticky="w")
        field = Entry(root, width=45)
        field.grid(row=row, column=1, padx=10, pady=3)
        fields.append(field)

    def clear():
        for field in fields:
            field.delete(0, END)

    def insert():
        values = [field.get() for field in fields]
        if all(value == "" for value in values):
            messagebox.showwarning("Empty input", "Please enter registration details.")
            return
        sheet.append(values)
        wb.save(excel_file)
        messagebox.showinfo("Saved", f"Saved to {excel_file.resolve()}")
        clear()
        fields[0].focus_set()

    Button(root, text="Submit", fg="Black", bg="Red", command=insert).grid(
        row=8, column=1, pady=10)
    root.mainloop()
    wb.save(excel_file)
    wb.close()


def program_36():
    from random import randint
    rand = randint(1, 100)
    counter = 0
    while True:
        try:
            num = int(input("Enter values between 1 and 100 (0 exit): "))
        except ValueError:
            print("Please enter an integer.")
            continue
        if num == 0:
            print("You canceled the game.")
            break
        elif num < rand:
            print("Enter a Higher Number.")
        elif num > rand:
            print("Enter a Lower Number.")
        else:
            print("Randomly selected number {0}!".format(rand))
            print("Your guess number {0}".format(counter + 1))
            break
        counter += 1


def program_37():
    favorite_fruits = ["apples", "erics", "oranges"]
    print("My favourite fruits {}".format(favorite_fruits))


def program_38():
    values = [3.14, 2.54, "CPU", "WINDOWS 10"]
    print(values)


def program_39():
    weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    print(weekdays[4])


def program_40():
    values = [3, 1, 2]
    values[0] = 1
    values[1] = 1
    values[2] = 2
    print(values)


def program_41():
    weekdays = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday"]
    print("Friday" in weekdays)
    print("Saturday" in weekdays)


def program_42():
    list1 = [10, 21, 4, 45, 66, 93]
    for num in list1:
        if num % 2 == 0:
            print(num, end=" ")
    print()


def program_43():
    numbers = [3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18]
    numbers2 = [21, 22, 23, 24, 25]
    combine = numbers + numbers2
    for number in combine:
        if number % 4 == 0:
            print(number)


def program_44():
    i = 1
    while i <= 12:
        print(" {}. class".format(i))
        i += 1


def program_45():
    number = input("enter a number: ")
    total = 0
    for digit in number:
        if digit.isdigit():
            total += int(digit)
    print("the sum of the digits of the number:", total)


def program_46():
    number = input("enter a number: ")
    multiplication = 1
    for digit in str(number):
        if digit.isdigit() and int(digit) != 0:
            multiplication *= int(digit)
    print("digits of the number: ", multiplication)


def program_47():
    text = input("enter a name: ")
    for i in range(10):
        print(text)


def program_48():
    text = input("enter a name: ")
    i = 0
    while i < 10:
        i += 1
        print(text)


def program_49():
    values = [35, 48, 52, 57, 68, 84]
    total = sum(values)
    piece = len(values)
    print(total / piece)


def program_50():
    total = 0
    while True:
        try:
            num = float(input("enter a number (0 to stop): "))
        except ValueError:
            print("Please enter a number.")
            continue
        if num == 0:
            break
        total += num
    print("The sum of the numbers: ", total)


PROGRAMS = {
    31: ("Read records from a MySQL database", program_31),
    32: ("Open a basic Tkinter window", program_32),
    33: ("Tkinter form entry", program_33),
    34: ("Tkinter ListBox", program_34),
    35: ("Registration form saved to Excel", program_35),
    36: ("Number guessing game", program_36),
    37: ("Favorite fruits list", program_37),
    38: ("Print a mixed list", program_38),
    39: ("Print a weekday by index", program_39),
    40: ("Change list values", program_40),
    41: ("Check weekdays in a list", program_41),
    42: ("Print even numbers in a list", program_42),
    43: ("Combine lists and print multiples of 4", program_43),
    44: ("Print 1 to 12 using while loop", program_44),
    45: ("Sum the digits of a number", program_45),
    46: ("Multiply nonzero digits", program_46),
    47: ("Print a name 10 times using for loop", program_47),
    48: ("Print a name 10 times using while loop", program_48),
    49: ("Find the mean using sum() and len()", program_49),
    50: ("Sum numbers until 0 is entered", program_50),
}


def main():
    while True:
        print("\n========== PLNG MIDTERM PRACTICE #3 ==========")
        for number, (description, _) in PROGRAMS.items():
            print(f"{number}. {description}")
        print("0. Exit")
        choice = input("Choose a program (31-50, or 0 to exit): ").strip()

        if choice == "0":
            print("Goodbye!")
            break
        if not choice.isdigit() or int(choice) not in PROGRAMS:
            print("Invalid choice. Please select a program from 31 to 50, or 0 to exit.")
            continue

        try:
            PROGRAMS[int(choice)][1]()
        except ModuleNotFoundError as error:
            print(f"Required module is missing: {error}. Install the needed package and try again.")
        except Exception as error:
            print(f"Program stopped because of an error: {error}")

        again = input("\nReturn to the menu? (Y/N): ").strip().upper()
        if again != "Y":
            print("Goodbye!")
            break


if __name__ == "__main__":
    main()
